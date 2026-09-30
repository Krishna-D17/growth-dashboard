import io
import csv
import json
import re
import uuid
from datetime import datetime, date
from enum import Enum
from typing import Dict, Any, List, Optional, Tuple, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.models.profile import Profile, ProfileSnapshot
from app.models.post import Post, PostSnapshot
from app.models.collection import CollectionJob
from app.models.anomaly import Anomaly
from app.services.analytics_service import analytics_service
from app.analytics.schemas import ProfileAnalyticsOverview


class ExportFormat(str, Enum):
    CSV = "csv"
    JSON = "json"
    XLSX = "xlsx"


def sanitize_filename_component(name: str) -> str:
    """Sanitizes username/filename component to prevent header injection and path traversal."""
    cleaned = re.sub(r"[^a-zA-Z0-9_\-]", "_", name)
    return cleaned if cleaned else "target"


def serialize_value(val: Any) -> Any:
    """Serializes special Python types to JSON/CSV primitive representations while preserving None."""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.isoformat()
    if isinstance(val, uuid.UUID):
        return str(val)
    if isinstance(val, Enum):
        return val.value
    if isinstance(val, (dict, list)):
        return json.dumps(val, default=str)
    return val


def build_csv_bytes(headers: List[str], rows: List[List[Any]]) -> bytes:
    """Renders tabular headers and rows into UTF-8 encoded CSV byte stream."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        serialized_row = [serialize_value(item) if item is not None else "" for item in row]
        writer.writerow(serialized_row)
    return output.getvalue().encode("utf-8")


def build_multi_section_csv_bytes(sections: Dict[str, Tuple[List[str], List[List[Any]]]]) -> bytes:
    """Renders multi-dataset sections into a single readable CSV byte stream."""
    output = io.StringIO()
    writer = csv.writer(output)
    for section_name, (headers, rows) in sections.items():
        writer.writerow([f"=== SECTION: {section_name.upper()} ==="])
        writer.writerow(headers)
        for row in rows:
            serialized_row = [serialize_value(item) if item is not None else "" for item in row]
            writer.writerow(serialized_row)
        writer.writerow([])  # Blank spacing line
    return output.getvalue().encode("utf-8")


def build_xlsx_bytes(sheets: Dict[str, Tuple[List[str], List[List[Any]]]]) -> bytes:
    """Generates an Excel (.xlsx) byte stream with styled headers and frozen panes using openpyxl."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default initial sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    for sheet_name, (headers, rows) in sheets.items():
        safe_title = sheet_name[:31]  # Excel max worksheet title limit
        ws = wb.create_sheet(title=safe_title)
        
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.freeze_panes = "A2"

        for row in rows:
            formatted_row = []
            for item in row:
                if item is None:
                    formatted_row.append(None)
                elif isinstance(item, (datetime, date)):
                    formatted_row.append(item.isoformat())
                elif isinstance(item, uuid.UUID):
                    formatted_row.append(str(item))
                elif isinstance(item, Enum):
                    formatted_row.append(item.value)
                elif isinstance(item, (dict, list)):
                    formatted_row.append(json.dumps(item, default=str))
                else:
                    formatted_row.append(item)
            ws.append(formatted_row)

        # Auto-adjust column width based on content
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value) if cell.value is not None else ""
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ExportService:
    """Service handling multi-format data exports (CSV, JSON, XLSX) for SocialScope profiles."""

    def _get_profile_or_404(self, db: Session, profile_id: uuid.UUID) -> Profile:
        profile = db.query(Profile).filter_by(id=profile_id).first()
        if not profile:
            raise ValueError(f"Profile with ID '{profile_id}' not found.")
        return profile

    # ------------------------------------------------------------------
    # 1. PROFILE EXPORT
    # ------------------------------------------------------------------
    def export_profile(
        self, db: Session, profile_id: uuid.UUID, fmt: ExportFormat
    ) -> Tuple[bytes, str, str]:
        profile = self._get_profile_or_404(db, profile_id)
        safe_name = sanitize_filename_component(profile.username)

        headers = [
            "id", "platform", "platform_profile_id", "username", "display_name",
            "profile_url", "bio", "verified", "collection_schedule",
            "last_collected_at", "created_at", "updated_at"
        ]
        row = [
            profile.id,
            profile.platform.value if hasattr(profile.platform, "value") else str(profile.platform),
            profile.platform_profile_id,
            profile.username,
            profile.display_name,
            profile.profile_url,
            profile.bio,
            profile.verified,
            profile.collection_schedule.value if hasattr(profile.collection_schedule, "value") else str(profile.collection_schedule),
            profile.last_collected_at,
            profile.created_at,
            profile.updated_at
        ]

        if fmt == ExportFormat.JSON:
            data_dict = {h: serialize_value(v) for h, v in zip(headers, row)}
            content = json.dumps(data_dict, indent=2).encode("utf-8")
            media_type = "application/json"
            filename = f"socialscope_{safe_name}_profile.json"
        elif fmt == ExportFormat.CSV:
            content = build_csv_bytes(headers, [row])
            media_type = "text/csv"
            filename = f"socialscope_{safe_name}_profile.csv"
        elif fmt == ExportFormat.XLSX:
            content = build_xlsx_bytes({"Profile": (headers, [row])})
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"socialscope_{safe_name}_profile.xlsx"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        return content, media_type, filename

    # ------------------------------------------------------------------
    # 2. PROFILE SNAPSHOTS EXPORT
    # ------------------------------------------------------------------
    def export_snapshots(
        self, db: Session, profile_id: uuid.UUID, fmt: ExportFormat
    ) -> Tuple[bytes, str, str]:
        profile = self._get_profile_or_404(db, profile_id)
        safe_name = sanitize_filename_component(profile.username)

        snapshots = db.query(ProfileSnapshot).filter_by(
            profile_id=profile_id
        ).order_by(ProfileSnapshot.collected_at.desc()).all()

        headers = [
            "id", "profile_id", "collected_at", "followers", "following",
            "post_count", "collector_version", "other_platform_metrics"
        ]
        rows = [
            [
                s.id, s.profile_id, s.collected_at, s.followers, s.following,
                s.post_count, s.collector_version, s.other_platform_metrics
            ]
            for s in snapshots
        ]

        if fmt == ExportFormat.JSON:
            list_data = [
                {h: serialize_value(v) for h, v in zip(headers, r)}
                for r in rows
            ]
            content = json.dumps(list_data, indent=2).encode("utf-8")
            media_type = "application/json"
            filename = f"socialscope_{safe_name}_snapshots.json"
        elif fmt == ExportFormat.CSV:
            content = build_csv_bytes(headers, rows)
            media_type = "text/csv"
            filename = f"socialscope_{safe_name}_snapshots.csv"
        elif fmt == ExportFormat.XLSX:
            content = build_xlsx_bytes({"Profile Snapshots": (headers, rows)})
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"socialscope_{safe_name}_snapshots.xlsx"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        return content, media_type, filename

    # ------------------------------------------------------------------
    # 3. POSTS & POST SNAPSHOTS EXPORT
    # ------------------------------------------------------------------
    def export_posts(
        self, db: Session, profile_id: uuid.UUID, fmt: ExportFormat
    ) -> Tuple[bytes, str, str]:
        profile = self._get_profile_or_404(db, profile_id)
        safe_name = sanitize_filename_component(profile.username)

        posts = db.query(Post).options(
            joinedload(Post.snapshots)
        ).filter_by(profile_id=profile_id).order_by(Post.posted_at.desc().nullslast()).all()

        headers = [
            "id", "profile_id", "platform_post_id", "url", "caption", "posted_at",
            "media_type", "created_at", "latest_likes", "latest_comments",
            "latest_shares", "latest_views", "latest_engagement"
        ]
        rows = []
        for p in posts:
            latest_snap = p.snapshots[0] if p.snapshots else None
            rows.append([
                p.id,
                p.profile_id,
                p.platform_post_id,
                p.url,
                p.caption,
                p.posted_at,
                p.media_type,
                p.created_at,
                latest_snap.likes if latest_snap else None,
                latest_snap.comments if latest_snap else None,
                latest_snap.shares if latest_snap else None,
                latest_snap.views if latest_snap else None,
                latest_snap.engagement if latest_snap else None,
            ])

        if fmt == ExportFormat.JSON:
            json_posts = []
            for p in posts:
                post_dict = {
                    "id": serialize_value(p.id),
                    "profile_id": serialize_value(p.profile_id),
                    "platform_post_id": p.platform_post_id,
                    "url": p.url,
                    "caption": p.caption,
                    "posted_at": serialize_value(p.posted_at),
                    "media_type": p.media_type,
                    "created_at": serialize_value(p.created_at),
                    "snapshots": [
                        {
                            "id": serialize_value(s.id),
                            "collected_at": serialize_value(s.collected_at),
                            "likes": s.likes,
                            "comments": s.comments,
                            "shares": s.shares,
                            "views": s.views,
                            "engagement": s.engagement,
                        }
                        for s in p.snapshots
                    ]
                }
                json_posts.append(post_dict)

            content = json.dumps(json_posts, indent=2).encode("utf-8")
            media_type = "application/json"
            filename = f"socialscope_{safe_name}_posts.json"

        elif fmt == ExportFormat.CSV:
            content = build_csv_bytes(headers, rows)
            media_type = "text/csv"
            filename = f"socialscope_{safe_name}_posts.csv"

        elif fmt == ExportFormat.XLSX:
            post_headers = headers
            post_rows = rows

            # Also prepare Post Snapshots sheet
            snap_headers = ["snapshot_id", "post_id", "collected_at", "likes", "comments", "shares", "views", "engagement"]
            snap_rows = []
            for p in posts:
                for s in p.snapshots:
                    snap_rows.append([s.id, p.id, s.collected_at, s.likes, s.comments, s.shares, s.views, s.engagement])

            sheets = {
                "Posts": (post_headers, post_rows),
                "Post Snapshots": (snap_headers, snap_rows)
            }
            content = build_xlsx_bytes(sheets)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"socialscope_{safe_name}_posts.xlsx"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        return content, media_type, filename

    # ------------------------------------------------------------------
    # 4. COLLECTION JOBS EXPORT
    # ------------------------------------------------------------------
    def export_jobs(
        self, db: Session, profile_id: uuid.UUID, fmt: ExportFormat
    ) -> Tuple[bytes, str, str]:
        profile = self._get_profile_or_404(db, profile_id)
        safe_name = sanitize_filename_component(profile.username)

        jobs = db.query(CollectionJob).filter_by(
            profile_id=profile_id
        ).order_by(CollectionJob.started_at.desc()).all()

        headers = [
            "id", "profile_id", "platform", "started_at", "completed_at",
            "status", "records_collected", "error_message", "collector_version"
        ]
        rows = [
            [
                j.id,
                j.profile_id,
                j.platform.value if hasattr(j.platform, "value") else str(j.platform),
                j.started_at,
                j.completed_at,
                j.status.value if hasattr(j.status, "value") else str(j.status),
                j.records_collected,
                j.error_message,
                j.collector_version,
            ]
            for j in jobs
        ]

        if fmt == ExportFormat.JSON:
            list_data = [
                {h: serialize_value(v) for h, v in zip(headers, r)}
                for r in rows
            ]
            content = json.dumps(list_data, indent=2).encode("utf-8")
            media_type = "application/json"
            filename = f"socialscope_{safe_name}_collection_jobs.json"
        elif fmt == ExportFormat.CSV:
            content = build_csv_bytes(headers, rows)
            media_type = "text/csv"
            filename = f"socialscope_{safe_name}_collection_jobs.csv"
        elif fmt == ExportFormat.XLSX:
            content = build_xlsx_bytes({"Collection Jobs": (headers, rows)})
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"socialscope_{safe_name}_collection_jobs.xlsx"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        return content, media_type, filename

    # ------------------------------------------------------------------
    # 5. ANALYTICS SUMMARY EXPORT
    # ------------------------------------------------------------------
    def export_analytics(
        self, db: Session, profile_id: uuid.UUID, fmt: ExportFormat, days: Optional[int] = None
    ) -> Tuple[bytes, str, str]:
        profile = self._get_profile_or_404(db, profile_id)
        safe_name = sanitize_filename_component(profile.username)

        overview = analytics_service.get_profile_analytics_overview(db, profile_id, days=days)

        if fmt == ExportFormat.JSON:
            content = json.dumps(overview.model_dump(), default=serialize_value, indent=2).encode("utf-8")
            media_type = "application/json"
            filename = f"socialscope_{safe_name}_analytics.json"

        elif fmt == ExportFormat.CSV:
            # Flatten analytics sections for tabular CSV presentation
            sections = {
                "Growth Analytics": (
                    ["metric", "value"],
                    [
                        ["current_followers", overview.growth.current_followers],
                        ["previous_followers", overview.growth.previous_followers],
                        ["absolute_growth", overview.growth.absolute_growth],
                        ["growth_percent", overview.growth.growth_percent],
                        ["growth_7d", overview.growth.growth_7d],
                        ["growth_percent_7d", overview.growth.growth_percent_7d],
                        ["growth_30d", overview.growth.growth_30d],
                        ["growth_percent_30d", overview.growth.growth_percent_30d],
                        ["growth_velocity", overview.growth.growth_velocity],
                        ["growth_acceleration", overview.growth.growth_acceleration],
                        ["start_date", overview.growth.start_date],
                        ["end_date", overview.growth.end_date],
                    ]
                ),
                "Engagement Analytics": (
                    ["metric", "value"],
                    [
                        ["total_engagement", overview.engagement.total_engagement],
                        ["average_engagement", overview.engagement.average_engagement],
                        ["median_engagement", overview.engagement.median_engagement],
                        ["engagement_rate", overview.engagement.engagement_rate],
                        ["available_metrics", ", ".join(overview.engagement.available_metrics)],
                        ["total_likes", overview.engagement.total_likes],
                        ["total_comments", overview.engagement.total_comments],
                        ["total_shares", overview.engagement.total_shares],
                        ["total_views", overview.engagement.total_views],
                        ["sample_post_count", overview.engagement.sample_post_count],
                    ]
                ),
                "Content Analytics": (
                    ["content_type", "post_count", "average_engagement", "median_engagement", "total_views", "average_views", "posting_frequency"],
                    [
                        [
                            c.content_type, c.post_count, c.average_engagement, c.median_engagement,
                            c.total_views, c.average_views, c.posting_frequency
                        ]
                        for c in overview.content.by_content_type
                    ]
                ),
                "Frequency Analytics": (
                    ["metric", "value"],
                    [
                        ["posts_per_day", overview.frequency.posts_per_day],
                        ["posts_per_week", overview.frequency.posts_per_week],
                        ["posts_per_month", overview.frequency.posts_per_month],
                        ["avg_posting_interval_hours", overview.frequency.avg_posting_interval_hours],
                        ["median_posting_interval_hours", overview.frequency.median_posting_interval_hours],
                        ["min_posting_interval_hours", overview.frequency.min_posting_interval_hours],
                        ["max_posting_interval_hours", overview.frequency.max_posting_interval_hours],
                    ]
                ),
                "Anomalies": (
                    ["detected_at", "metric", "observed_value", "severity", "method", "description"],
                    [
                        [a.detected_at, a.metric, a.observed_value, a.severity, a.method, a.description]
                        for a in overview.anomalies
                    ]
                )
            }
            content = build_multi_section_csv_bytes(sections)
            media_type = "text/csv"
            filename = f"socialscope_{safe_name}_analytics.csv"

        elif fmt == ExportFormat.XLSX:
            sheets = {
                "Growth": (
                    ["Metric", "Value"],
                    [
                        ["Current Followers", overview.growth.current_followers],
                        ["Previous Followers", overview.growth.previous_followers],
                        ["Absolute Growth", overview.growth.absolute_growth],
                        ["Growth %", overview.growth.growth_percent],
                        ["7-Day Growth", overview.growth.growth_7d],
                        ["7-Day Growth %", overview.growth.growth_percent_7d],
                        ["30-Day Growth", overview.growth.growth_30d],
                        ["30-Day Growth %", overview.growth.growth_percent_30d],
                        ["Growth Velocity (Followers/Day)", overview.growth.growth_velocity],
                        ["Growth Acceleration", overview.growth.growth_acceleration],
                    ]
                ),
                "Engagement": (
                    ["Metric", "Value"],
                    [
                        ["Total Engagement", overview.engagement.total_engagement],
                        ["Average Engagement", overview.engagement.average_engagement],
                        ["Median Engagement", overview.engagement.median_engagement],
                        ["Engagement Rate", overview.engagement.engagement_rate],
                        ["Available Metrics", ", ".join(overview.engagement.available_metrics)],
                        ["Total Likes", overview.engagement.total_likes],
                        ["Total Comments", overview.engagement.total_comments],
                        ["Total Shares", overview.engagement.total_shares],
                        ["Total Views", overview.engagement.total_views],
                        ["Sample Post Count", overview.engagement.sample_post_count],
                    ]
                ),
                "Content": (
                    ["Content Type", "Post Count", "Avg Engagement", "Median Engagement", "Total Views", "Avg Views", "Posts / Day"],
                    [
                        [
                            c.content_type, c.post_count, c.average_engagement, c.median_engagement,
                            c.total_views, c.average_views, c.posting_frequency
                        ]
                        for c in overview.content.by_content_type
                    ]
                ),
                "Frequency": (
                    ["Metric", "Value"],
                    [
                        ["Posts / Day", overview.frequency.posts_per_day],
                        ["Posts / Week", overview.frequency.posts_per_week],
                        ["Posts / Month", overview.frequency.posts_per_month],
                        ["Avg Posting Interval (hrs)", overview.frequency.avg_posting_interval_hours],
                        ["Median Posting Interval (hrs)", overview.frequency.median_posting_interval_hours],
                        ["Min Posting Interval (hrs)", overview.frequency.min_posting_interval_hours],
                        ["Max Posting Interval (hrs)", overview.frequency.max_posting_interval_hours],
                    ]
                ),
                "Anomalies": (
                    ["Detected At", "Metric", "Observed Value", "Severity", "Method", "Description"],
                    [
                        [a.detected_at, a.metric, a.observed_value, a.severity, a.method, a.description]
                        for a in overview.anomalies
                    ]
                )
            }
            content = build_xlsx_bytes(sheets)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"socialscope_{safe_name}_analytics.xlsx"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        return content, media_type, filename

    # ------------------------------------------------------------------
    # 6. COMPLETE MULTI-DATASET REPORT EXPORT
    # ------------------------------------------------------------------
    def export_complete_report(
        self, db: Session, profile_id: uuid.UUID, fmt: ExportFormat, days: Optional[int] = None
    ) -> Tuple[bytes, str, str]:
        profile = self._get_profile_or_404(db, profile_id)
        safe_name = sanitize_filename_component(profile.username)

        # Gather profile, snapshots, posts, jobs, and analytics overview
        snapshots = db.query(ProfileSnapshot).filter_by(profile_id=profile_id).order_by(ProfileSnapshot.collected_at.desc()).all()
        posts = db.query(Post).options(joinedload(Post.snapshots)).filter_by(profile_id=profile_id).order_by(Post.posted_at.desc().nullslast()).all()
        jobs = db.query(CollectionJob).filter_by(profile_id=profile_id).order_by(CollectionJob.started_at.desc()).all()
        overview = analytics_service.get_profile_analytics_overview(db, profile_id, days=days)

        if fmt == ExportFormat.JSON:
            report_data = {
                "profile": {
                    "id": str(profile.id),
                    "platform": profile.platform.value if hasattr(profile.platform, "value") else str(profile.platform),
                    "platform_profile_id": profile.platform_profile_id,
                    "username": profile.username,
                    "display_name": profile.display_name,
                    "profile_url": profile.profile_url,
                    "bio": profile.bio,
                    "verified": profile.verified,
                    "collection_schedule": profile.collection_schedule.value if hasattr(profile.collection_schedule, "value") else str(profile.collection_schedule),
                    "last_collected_at": serialize_value(profile.last_collected_at),
                    "created_at": serialize_value(profile.created_at),
                    "updated_at": serialize_value(profile.updated_at),
                },
                "snapshots": [
                    {
                        "id": str(s.id),
                        "collected_at": serialize_value(s.collected_at),
                        "followers": s.followers,
                        "following": s.following,
                        "post_count": s.post_count,
                        "collector_version": s.collector_version,
                        "other_platform_metrics": s.other_platform_metrics,
                    }
                    for s in snapshots
                ],
                "posts": [
                    {
                        "id": str(p.id),
                        "platform_post_id": p.platform_post_id,
                        "url": p.url,
                        "caption": p.caption,
                        "posted_at": serialize_value(p.posted_at),
                        "media_type": p.media_type,
                        "created_at": serialize_value(p.created_at),
                        "snapshots": [
                            {
                                "id": str(s.id),
                                "collected_at": serialize_value(s.collected_at),
                                "likes": s.likes,
                                "comments": s.comments,
                                "shares": s.shares,
                                "views": s.views,
                                "engagement": s.engagement,
                            }
                            for s in p.snapshots
                        ]
                    }
                    for p in posts
                ],
                "collection_jobs": [
                    {
                        "id": str(j.id),
                        "started_at": serialize_value(j.started_at),
                        "completed_at": serialize_value(j.completed_at),
                        "status": j.status.value if hasattr(j.status, "value") else str(j.status),
                        "records_collected": j.records_collected,
                        "error_message": j.error_message,
                    }
                    for j in jobs
                ],
                "analytics": overview.model_dump()
            }
            content = json.dumps(report_data, default=serialize_value, indent=2).encode("utf-8")
            media_type = "application/json"
            filename = f"socialscope_{safe_name}_report.json"

        elif fmt == ExportFormat.CSV:
            # Generate multi-section CSV for complete report
            sections = {
                "Profile": (
                    ["username", "platform", "followers", "verified", "last_collected_at"],
                    [[profile.username, profile.platform.value if hasattr(profile.platform, "value") else str(profile.platform), overview.growth.current_followers, profile.verified, profile.last_collected_at]]
                ),
                "Profile Snapshots": (
                    ["collected_at", "followers", "following", "post_count"],
                    [[s.collected_at, s.followers, s.following, s.post_count] for s in snapshots]
                ),
                "Posts": (
                    ["platform_post_id", "url", "posted_at", "media_type", "likes", "comments", "shares", "views", "engagement"],
                    [
                        [
                            p.platform_post_id, p.url, p.posted_at, p.media_type,
                            p.snapshots[0].likes if p.snapshots else None,
                            p.snapshots[0].comments if p.snapshots else None,
                            p.snapshots[0].shares if p.snapshots else None,
                            p.snapshots[0].views if p.snapshots else None,
                            p.snapshots[0].engagement if p.snapshots else None,
                        ]
                        for p in posts
                    ]
                ),
                "Collection Jobs": (
                    ["started_at", "completed_at", "status", "records_collected", "error_message"],
                    [[j.started_at, j.completed_at, j.status.value if hasattr(j.status, "value") else str(j.status), j.records_collected, j.error_message] for j in jobs]
                ),
                "Growth Summary": (
                    ["current_followers", "growth_7d", "growth_30d", "velocity"],
                    [[overview.growth.current_followers, overview.growth.growth_7d, overview.growth.growth_30d, overview.growth.growth_velocity]]
                )
            }
            content = build_multi_section_csv_bytes(sections)
            media_type = "text/csv"
            filename = f"socialscope_{safe_name}_report.csv"

        elif fmt == ExportFormat.XLSX:
            # Build individual worksheets for complete workbook
            profile_headers = [
                "id", "platform", "platform_profile_id", "username", "display_name",
                "profile_url", "bio", "verified", "collection_schedule", "last_collected_at", "created_at"
            ]
            profile_row = [[
                profile.id,
                profile.platform.value if hasattr(profile.platform, "value") else str(profile.platform),
                profile.platform_profile_id,
                profile.username,
                profile.display_name,
                profile.profile_url,
                profile.bio,
                profile.verified,
                profile.collection_schedule.value if hasattr(profile.collection_schedule, "value") else str(profile.collection_schedule),
                profile.last_collected_at,
                profile.created_at
            ]]

            snap_headers = ["id", "collected_at", "followers", "following", "post_count", "collector_version"]
            snap_rows = [[s.id, s.collected_at, s.followers, s.following, s.post_count, s.collector_version] for s in snapshots]

            post_headers = ["id", "platform_post_id", "url", "caption", "posted_at", "media_type", "latest_likes", "latest_comments", "latest_shares", "latest_views", "latest_engagement"]
            post_rows = [
                [
                    p.id, p.platform_post_id, p.url, p.caption, p.posted_at, p.media_type,
                    p.snapshots[0].likes if p.snapshots else None,
                    p.snapshots[0].comments if p.snapshots else None,
                    p.snapshots[0].shares if p.snapshots else None,
                    p.snapshots[0].views if p.snapshots else None,
                    p.snapshots[0].engagement if p.snapshots else None,
                ]
                for p in posts
            ]

            post_snap_headers = ["snapshot_id", "post_id", "collected_at", "likes", "comments", "shares", "views", "engagement"]
            post_snap_rows = []
            for p in posts:
                for ps in p.snapshots:
                    post_snap_rows.append([ps.id, p.id, ps.collected_at, ps.likes, ps.comments, ps.shares, ps.views, ps.engagement])

            job_headers = ["id", "started_at", "completed_at", "status", "records_collected", "error_message"]
            job_rows = [[j.id, j.started_at, j.completed_at, j.status.value if hasattr(j.status, "value") else str(j.status), j.records_collected, j.error_message] for j in jobs]

            sheets = {
                "Profile": (profile_headers, profile_row),
                "Profile Snapshots": (snap_headers, snap_rows),
                "Posts": (post_headers, post_rows),
                "Post Snapshots": (post_snap_headers, post_snap_rows),
                "Collection Jobs": (job_headers, job_rows),
                "Growth": (
                    ["Metric", "Value"],
                    [
                        ["Current Followers", overview.growth.current_followers],
                        ["Previous Followers", overview.growth.previous_followers],
                        ["Absolute Growth", overview.growth.absolute_growth],
                        ["Growth %", overview.growth.growth_percent],
                        ["7-Day Growth", overview.growth.growth_7d],
                        ["30-Day Growth", overview.growth.growth_30d],
                        ["Growth Velocity", overview.growth.growth_velocity],
                        ["Growth Acceleration", overview.growth.growth_acceleration],
                    ]
                ),
                "Engagement": (
                    ["Metric", "Value"],
                    [
                        ["Total Engagement", overview.engagement.total_engagement],
                        ["Average Engagement", overview.engagement.average_engagement],
                        ["Median Engagement", overview.engagement.median_engagement],
                        ["Engagement Rate", overview.engagement.engagement_rate],
                        ["Total Likes", overview.engagement.total_likes],
                        ["Total Comments", overview.engagement.total_comments],
                        ["Total Shares", overview.engagement.total_shares],
                        ["Total Views", overview.engagement.total_views],
                    ]
                ),
                "Content": (
                    ["Content Type", "Post Count", "Avg Engagement", "Median Engagement", "Total Views", "Avg Views"],
                    [
                        [c.content_type, c.post_count, c.average_engagement, c.median_engagement, c.total_views, c.average_views]
                        for c in overview.content.by_content_type
                    ]
                ),
                "Frequency": (
                    ["Metric", "Value"],
                    [
                        ["Posts / Day", overview.frequency.posts_per_day],
                        ["Posts / Week", overview.frequency.posts_per_week],
                        ["Posts / Month", overview.frequency.posts_per_month],
                        ["Avg Posting Interval (hrs)", overview.frequency.avg_posting_interval_hours],
                    ]
                ),
                "Anomalies": (
                    ["Detected At", "Metric", "Observed Value", "Severity", "Description"],
                    [
                        [a.detected_at, a.metric, a.observed_value, a.severity, a.description]
                        for a in overview.anomalies
                    ]
                )
            }
            content = build_xlsx_bytes(sheets)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"socialscope_{safe_name}_report.xlsx"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        return content, media_type, filename


export_service = ExportService()
