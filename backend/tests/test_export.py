import io
import uuid
import pytest
import openpyxl
from datetime import datetime, timezone, timedelta
from fastapi.testclient import TestClient

from app.main import app
from app.database.session import SessionLocal
from app.models.enums import SocialPlatform, JobStatus
from app.models.profile import Profile, ProfileSnapshot
from app.models.post import Post, PostSnapshot
from app.models.collection import CollectionJob

client = TestClient(app)


@pytest.fixture
def db():
    session = SessionLocal()
    session.query(Profile).delete()
    session.commit()
    try:
        yield session
    finally:
        session.rollback()
        session.query(Profile).delete()
        session.commit()
        session.close()


def create_export_test_data(db, username="export_user"):
    now = datetime.now(timezone.utc)
    profile = Profile(
        platform=SocialPlatform.INSTAGRAM,
        username=username,
        display_name="Export User Display",
        profile_url=f"https://instagram.com/{username}",
        bio="Test Bio",
        verified=True
    )
    db.add(profile)
    db.commit()
    db.refresh(profile)

    # Snapshots
    s1 = ProfileSnapshot(
        profile_id=profile.id,
        collected_at=now - timedelta(days=14),
        followers=1000,
        following=500,
        post_count=10
    )
    s2 = ProfileSnapshot(
        profile_id=profile.id,
        collected_at=now - timedelta(days=7),
        followers=1200,
        following=510,
        post_count=12
    )
    db.add_all([s1, s2])

    # Posts & Snapshots
    p1 = Post(
        profile_id=profile.id,
        platform_post_id="post_101",
        url="https://instagram.com/p/post_101",
        caption="Sample Post 101",
        posted_at=now - timedelta(days=5),
        media_type="IMAGE"
    )
    db.add(p1)
    db.commit()
    db.refresh(p1)

    ps1 = PostSnapshot(
        post_id=p1.id,
        collected_at=now - timedelta(days=5),
        likes=150,
        comments=20,
        shares=None,  # Null metric preservation
        views=None,   # Null metric preservation
        engagement=170.0
    )
    db.add(ps1)

    # Collection Job
    job = CollectionJob(
        profile_id=profile.id,
        platform=SocialPlatform.INSTAGRAM,
        started_at=now - timedelta(days=1),
        completed_at=now - timedelta(days=1, minutes=-5),
        status=JobStatus.SUCCESS,
        records_collected=5
    )
    db.add(job)
    db.commit()

    return profile


def test_export_profile_csv(db):
    profile = create_export_test_data(db, "csv_user")
    res = client.get(f"/api/profiles/{profile.id}/export/profile?format=csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert 'attachment; filename="socialscope_csv_user_profile.csv"' in res.headers["content-disposition"]
    body = res.text
    assert "id,platform,platform_profile_id,username" in body
    assert "csv_user" in body
    assert "instagram" in body


def test_export_profile_json(db):
    profile = create_export_test_data(db, "json_user")
    res = client.get(f"/api/profiles/{profile.id}/export/profile?format=json")
    assert res.status_code == 200
    assert "application/json" in res.headers["content-type"]
    assert 'attachment; filename="socialscope_json_user_profile.json"' in res.headers["content-disposition"]
    data = res.json()
    assert data["username"] == "json_user"
    assert data["platform"] == "instagram"
    assert data["verified"] is True


def test_export_profile_xlsx(db):
    profile = create_export_test_data(db, "xlsx_user")
    res = client.get(f"/api/profiles/{profile.id}/export/profile?format=xlsx")
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert 'attachment; filename="socialscope_xlsx_user_profile.xlsx"' in res.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Profile" in wb.sheetnames
    ws = wb["Profile"]
    assert ws.cell(row=1, column=4).value == "username"
    assert ws.cell(row=2, column=4).value == "xlsx_user"


def test_export_snapshots(db):
    profile = create_export_test_data(db, "snap_user")
    res = client.get(f"/api/profiles/{profile.id}/export/snapshots?format=json")
    assert res.status_code == 200
    data = res.json()
    assert len(data) == 2
    assert data[0]["followers"] in [1000, 1200]


def test_export_posts(db):
    profile = create_export_test_data(db, "posts_user")
    res = client.get(f"/api/profiles/{profile.id}/export/posts?format=json")
    assert res.status_code == 200
    data = res.json()
    assert len(data) == 1
    assert data[0]["platform_post_id"] == "post_101"
    assert data[0]["snapshots"][0]["likes"] == 150
    # Verify null preservation in JSON
    assert data[0]["snapshots"][0]["shares"] is None


def test_export_jobs(db):
    profile = create_export_test_data(db, "jobs_user")
    res = client.get(f"/api/profiles/{profile.id}/export/jobs?format=json")
    assert res.status_code == 200
    data = res.json()
    assert len(data) == 1
    assert data[0]["status"] == "success"
    assert data[0]["records_collected"] == 5


def test_export_analytics(db):
    profile = create_export_test_data(db, "analytics_user")
    res = client.get(f"/api/profiles/{profile.id}/export/analytics?format=json&days=30")
    assert res.status_code == 200
    data = res.json()
    assert "growth" in data
    assert "engagement" in data
    assert data["growth"]["current_followers"] == 1200


def test_export_complete_report_xlsx(db):
    profile = create_export_test_data(db, "report_user")
    res = client.get(f"/api/profiles/{profile.id}/export/report?format=xlsx")
    assert res.status_code == 200
    assert 'attachment; filename="socialscope_report_user_report.xlsx"' in res.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    expected_sheets = ["Profile", "Profile Snapshots", "Posts", "Post Snapshots", "Collection Jobs", "Growth", "Engagement", "Content", "Frequency", "Anomalies"]
    for s in expected_sheets:
        assert s in wb.sheetnames

    # Check representative sheet contents
    ws_prof = wb["Profile"]
    assert ws_prof.cell(row=2, column=4).value == "report_user"

    ws_posts = wb["Posts"]
    assert ws_posts.cell(row=2, column=2).value == "post_101"


def test_export_nonexistent_profile_404(db):
    fake_id = uuid.uuid4()
    res = client.get(f"/api/profiles/{fake_id}/export/profile?format=csv")
    assert res.status_code == 404
    assert "not found" in res.json()["detail"].lower()


def test_export_invalid_format_422(db):
    profile = create_export_test_data(db, "fmt_user")
    res = client.get(f"/api/profiles/{profile.id}/export/profile?format=invalid_fmt")
    assert res.status_code == 422


def test_export_null_preservation(db):
    profile = create_export_test_data(db, "null_user")
    res = client.get(f"/api/profiles/{profile.id}/export/posts?format=json")
    assert res.status_code == 200
    post = res.json()[0]
    snap = post["snapshots"][0]
    assert snap["likes"] == 150
    assert snap["shares"] is None
    assert snap["views"] is None
